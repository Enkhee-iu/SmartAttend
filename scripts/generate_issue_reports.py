#!/usr/bin/env python3
r"""Generate Hitachi City help desk and maintenance monthly reports.

Double-clicking the Windows EXE opens a file picker for issues.csv. The two
Excel files are written to the same folder as the selected CSV file.
"""

from __future__ import annotations

import argparse
import calendar
import csv
import io
import re
import sys
import warnings
from copy import copy
from datetime import date, datetime
from pathlib import Path

warnings.filterwarnings(
    "ignore",
    message="Title is more than 31 characters.*",
    category=UserWarning,
    module="openpyxl.workbook.child",
)

try:
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
    from openpyxl.utils import get_column_letter
except ModuleNotFoundError as exc:
    raise SystemExit(
        "openpyxl is required. Install it with: python -m pip install -r requirements.txt"
    ) from exc


ENCODINGS_TO_TRY = ("utf-8-sig", "cp932", "shift_jis", "utf-8")

HELP_DESK_HEADERS = [
    "項番",
    "受付No",
    "受付日",
    "作業完了日",
    "状況",
    "部署",
    "依頼者",
    "依頼内容（問い合わせ含む）",
    "対応内容",
]

MAINTENANCE_HEADERS = [
    "項番",
    "受付No",
    "受付日",
    "作業完了日",
    "状況",
    "部署",
    "ご担当者",
    "機器名",
    "修理品型式",
    "製造番号",
    "症状",
    "原因",
    "処置",
]

HELPDESK_MARKER = "ヘルプデスク"
MAINTENANCE_MARKER = "賃貸借契約"
MISSING_MARKER_FLAG = "missing_marker"
SPLIT_MARKER_FLAG = "split_marker"

def parse_args() -> argparse.Namespace:
    parser = argparse.ArgumentParser(
        description="Read issues.csv and create help desk / maintenance Excel reports.",
    )
    parser.add_argument(
        "issues_csv",
        nargs="?",
        default=None,
        help="Path to issues.csv. If omitted, a file picker is opened.",
    )
    parser.add_argument(
        "--output-dir",
        default=None,
        help="Directory where the two Excel files are written. Default: same folder as issues.csv",
    )
    parser.add_argument(
        "--gui",
        action="store_true",
        help="Open a file picker for issues.csv and save reports next to the selected file.",
    )
    return parser.parse_args()


def read_csv_rows(path: Path) -> tuple[list[str], list[list[str]], str]:
    raw = path.read_bytes()
    decoded_candidates: list[tuple[int, str, str]] = []

    for encoding in ENCODINGS_TO_TRY:
        try:
            text = raw.decode(encoding)
        except UnicodeDecodeError:
            continue
        decoded_candidates.append((text.count("\ufffd"), encoding, text))

    if decoded_candidates:
        _, encoding, text = min(decoded_candidates, key=lambda item: item[0])
    else:
        encoding = "utf-8-replace"
        text = raw.decode("utf-8", errors="replace")

    reader = csv.reader(io.StringIO(text, newline=""))
    rows = [row for row in reader if any(cell.strip() for cell in row)]
    if not rows:
        raise ValueError(f"No rows found in CSV: {path}")

    max_columns = max(len(row) for row in rows)
    padded_rows = [row + [""] * (max_columns - len(row)) for row in rows]
    header = padded_rows[0]
    data_rows = padded_rows[1:]
    return header, data_rows, encoding


def fullwidth(value: int) -> str:
    return str(value).translate(str.maketrans("0123456789", "０１２３４５６７８９"))


def previous_month(today: date | None = None) -> tuple[int, int, int]:
    today = today or date.today()
    year = today.year
    month = today.month - 1
    if month == 0:
        month = 12
        year -= 1
    return year, month, calendar.monthrange(year, month)[1]


def report_filenames(year: int, month: int) -> tuple[str, str]:
    return (
        f"1.日立市様ヘルプデスク{year}年{month}月分ご報告.xlsx",
        f"2.日立市様各種賃貸借契約に対する保守対応定期報告書{year}年{month}月分ご報告.xlsx",
    )


def rows_as_dicts(header: list[str], rows: list[list[str]]) -> list[dict[str, str]]:
    return [dict(zip(header, row)) for row in rows]


def value(row: dict[str, str], key: str) -> str:
    return (row.get(key, "") or "").strip()


def filled(text: str) -> str:
    text = (text or "").strip()
    return text if text else "ー"


def parse_date(text: str):
    text = (text or "").strip()
    if not text:
        return "ー"
    for fmt in ("%Y/%m/%d", "%Y-%m-%d"):
        try:
            return datetime.strptime(text, fmt)
        except ValueError:
            pass
    return text


def sort_key(row: dict[str, str]) -> tuple[datetime, datetime, str]:
    received_date = parse_date(value(row, "受付日"))
    if not isinstance(received_date, datetime):
        received_date = datetime.min
    completed_date = parse_date(value(row, "作業完了日"))
    if not isinstance(completed_date, datetime):
        completed_date = datetime.min
    return received_date, completed_date, ticket_number(value(row, "題名"))


def ticket_number(title: str) -> str:
    match = re.search(r"\[(HB\d+)\]", title or "")
    return match.group(1) if match else filled(title)


def department_from_title(title: str) -> str:
    text = re.sub(r"\[HB\d+\]", "", title or "")
    text = re.sub(r"【[^】]*】", "", text)
    text = text.replace("日立市役所", "")
    text = text.replace("日立市", "")
    text = text.replace("日立保健センター", "保健センター")
    text = re.sub(r"[\[【].*?[\]】]", "", text)
    text = text.replace("　", " ")
    text = re.sub(r"\s+", " ", text).strip()
    return filled(text.strip(" -_　"))


def department(row: dict[str, str]) -> str:
    return value(row, "部署") or department_from_title(value(row, "題名"))


def combined_text(row: dict[str, str]) -> str:
    keys = ("症状", "原因", "処置", "社内メモ", "連絡欄", "最新のコメント")
    return "\n".join(value(row, key) for key in keys if value(row, key))


def split_work_items(row: dict[str, str]) -> list[tuple[dict[str, str], str, str]]:
    """Classify work items from the markers in 連絡欄.

    The markers are authoritative:
    - ヘルプデスク only -> help desk
    - 賃貸借契約 only -> maintenance
    - both -> split into one row for each report
    - neither -> maintenance with a red-font warning marker
    """
    contact = value(row, "連絡欄")
    has_helpdesk_marker = HELPDESK_MARKER in contact
    has_maintenance_marker = MAINTENANCE_MARKER in contact

    if has_helpdesk_marker and has_maintenance_marker:
        return [(row.copy(), "maintenance", SPLIT_MARKER_FLAG), (row.copy(), "helpdesk", SPLIT_MARKER_FLAG)]
    if has_helpdesk_marker:
        return [(row, "helpdesk", "")]
    if has_maintenance_marker:
        return [(row, "maintenance", "")]
    if not has_helpdesk_marker and not has_maintenance_marker:
        return [(row, "maintenance", MISSING_MARKER_FLAG)]


def helpdesk_row(row: dict[str, str], number: int) -> list[object]:
    return [
        number,
        ticket_number(value(row, "題名")),
        parse_date(value(row, "受付日")),
        parse_date(value(row, "作業完了日")),
        filled(value(row, "状況")),
        department(row),
        filled(value(row, "客先担当")),
        filled(value(row, "症状")),
        filled(value(row, "処置")),
    ]


def maintenance_row(row: dict[str, str], number: int) -> list[object]:
    return [
        number,
        ticket_number(value(row, "題名")),
        parse_date(value(row, "受付日")),
        parse_date(value(row, "作業完了日")),
        filled(value(row, "状況")),
        department(row),
        filled(value(row, "客先担当")),
        filled(value(row, "ホスト")),
        filled(value(row, "修理品型式")),
        filled(value(row, "製造番号")),
        filled(value(row, "症状")),
        filled(value(row, "原因")),
        filled(value(row, "処置")),
    ]


def apply_cell_style(cell, *, fill=None, font=None, border=None, alignment=None, number_format=None):
    if fill is not None:
        cell.fill = fill
    if font is not None:
        cell.font = font
    if border is not None:
        cell.border = border
    if alignment is not None:
        cell.alignment = alignment
    if number_format is not None:
        cell.number_format = number_format


def style_report_sheet(
    ws,
    widths: list[int],
    header_row: int = 1,
    start_col: int = 1,
    data_start_row: int | None = None,
    left_aligned_columns: set[int] | None = None,
    border_style: str = "thin",
) -> None:
    data_start_row = data_start_row or header_row + 1
    side = Side(style=border_style, color="000000")
    border = Border(left=side, right=side, top=side, bottom=side)
    header_font = Font(name="ＭＳ Ｐゴシック", bold=True, size=10)
    body_font = Font(name="ＭＳ Ｐゴシック", size=10)
    left_aligned_columns = left_aligned_columns or set()
    for row in ws.iter_rows(min_row=header_row, max_row=ws.max_row, min_col=start_col, max_col=start_col + len(widths) - 1):
        for cell in row:
            is_header = cell.row == header_row
            horizontal = "center" if is_header or cell.column not in left_aligned_columns else "left"
            apply_cell_style(
                cell,
                font=header_font if is_header else body_font,
                border=border,
                alignment=Alignment(horizontal=horizontal, vertical="center", wrap_text=True),
            )
            if isinstance(cell.value, datetime):
                cell.number_format = "yyyy/m/d"
    for offset, width in enumerate(widths):
        ws.column_dimensions[get_column_letter(start_col + offset)].width = width
    ws.freeze_panes = ws.cell(data_start_row, start_col).coordinate


def mark_output_row(ws, row_number: int, start_col: int, column_count: int, flag: str) -> None:
    if not flag:
        return
    split_fill = PatternFill("solid", fgColor="D9D9D9")
    red_font = Font(name="ＭＳ Ｐゴシック", size=10, color="FF0000")
    for col in range(start_col, start_col + column_count):
        cell = ws.cell(row_number, col)
        if flag == SPLIT_MARKER_FLAG:
            cell.fill = copy(split_fill)
        elif flag == MISSING_MARKER_FLAG:
            cell.font = copy(red_font)


def setup_summary_sheet(ws, title: str, period: str, count: int, remarks: str, is_maintenance: bool) -> None:
    ws.title = "特記事項"
    ws.sheet_view.showGridLines = False
    for col in range(1, 12):
        ws.column_dimensions[get_column_letter(col)].width = 14
    ws.column_dimensions["A"].width = 12
    for row in range(1, 31):
        ws.row_dimensions[row].height = 20
    yellow = PatternFill("solid", fgColor="FFFF00")
    green = PatternFill("solid", fgColor="C6EFCE")
    thin = Side(style="thin", color="000000")
    dotted = Side(style="dotted", color="808080")
    border = Border(left=thin, right=thin, top=thin, bottom=thin)
    ws.merge_cells("A3:K3")
    ws["A3"] = title
    ws["A3"].fill = yellow
    ws["A3"].font = Font(name="ＭＳ Ｐゴシック", bold=True, size=12)
    ws["A3"].alignment = Alignment(horizontal="center", vertical="center")
    rows = [
        (4, "保守期間", period),
        (5, "報    告", "エイチ・シー・ネットワークス株式会社"),
        (6, "責 任 者", "宮本　良一"),
        (7, "担    当", "保守：池田 修司 　営業：小瀬 真弘"),
    ]
    for row_num, label, text in rows:
        ws.merge_cells(start_row=row_num, start_column=1, end_row=row_num, end_column=3)
        ws.merge_cells(start_row=row_num, start_column=4, end_row=row_num, end_column=11)
        ws.cell(row_num, 1, label).alignment = Alignment(horizontal="center", vertical="center")
        ws.cell(row_num, 4, text)
    ws.merge_cells("A8:K8")
    ws["A8"] = "特記事項"
    ws["A8"].fill = green
    ws["A8"].alignment = Alignment(horizontal="center")
    sections = [
        (9, 13, "トピックス", ["特にございません。"]),
        (
            14,
            19,
            "対応状況",
            [
                f"今月は、{fullwidth(count)}件{'作業対応' if is_maintenance else '受付対応'}になります。",
                f"詳細は次紙によります。（{'作業完了' if is_maintenance else '完了'}：{fullwidth(count)}件）",
                f"（内訳は、ハード修理：{fullwidth(count)}件となっております。）" if is_maintenance else "",
            ],
        ),
        (20, 25, "今月の\n問題点", ["特にございません。"]),
        (26, 29, "備考", [remarks]),
    ]
    for start, end, label, lines in sections:
        ws.merge_cells(start_row=start, start_column=1, end_row=end, end_column=1)
        ws.cell(start, 1, label).alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        for idx, line in enumerate(line for line in lines if line):
            ws.merge_cells(start_row=start + idx, start_column=2, end_row=start + idx, end_column=11)
            ws.cell(start + idx, 2, line).alignment = Alignment(vertical="top", wrap_text=True)
        for row in range(start, end + 1):
            for col in range(2, 12):
                ws.cell(row, col).border = Border(bottom=dotted)
    for row in ws.iter_rows(min_row=3, max_row=29, min_col=1, max_col=11):
        for cell in row:
            if cell.row <= 8 or cell.column == 1:
                cell.border = border
            cell.font = Font(name="ＭＳ Ｐゴシック", size=10, bold=cell.row in (3, 8))


def maintenance_remarks(rows: list[tuple[dict[str, str], str]]) -> str:
    mouse = sum(1 for row, _flag in rows if "マウス" in combined_text(row))
    base = sum(1 for row, _flag in rows if "ベースエンクロージャ" in combined_text(row) or "ベースエンクロージャー" in combined_text(row))
    ac = 0
    for row, _flag in rows:
        action_text = "\n".join(value(row, key) for key in ("処置", "連絡欄", "社内メモ"))
        if ("ACアダプタ" in action_text or "アダプタ" in action_text) and "マザーボード" not in action_text:
            ac += 1
    return f"・マウス交換：{fullwidth(mouse)}件、ベースエンクロージャー交換：{fullwidth(base)}件、ACアダプタ交換：{fullwidth(ac)}件ございました。"


def write_helpdesk_workbook(path: Path, rows: list[tuple[dict[str, str], str]], year: int, month: int, last_day: int) -> None:
    wb = Workbook()
    summary = wb.active
    title = f"日立市役所様 本庁及び出先機関ヘルプデスク保守サービス＜{fullwidth(year)}年{fullwidth(month)}月度＞月次報告書"
    period = f"{fullwidth(year)}年{fullwidth(month)}月1日～{fullwidth(year)}年{fullwidth(month)}月{fullwidth(last_day)}日"
    setup_summary_sheet(summary, title, period, len(rows), "特にございません。", False)
    ws = wb.create_sheet("データ")
    ws.append(HELP_DESK_HEADERS)
    for index, (row, _flag) in enumerate(rows, start=1):
        ws.append(helpdesk_row(row, index))
    style_report_sheet(
        ws,
        [8, 12, 13, 13, 12, 24, 18, 50, 60],
        left_aligned_columns={8, 9},
        border_style="medium",
    )
    for row_number, (_row, flag) in enumerate(rows, start=2):
        mark_output_row(ws, row_number, 1, len(HELP_DESK_HEADERS), flag)
    wb.save(path)


def write_maintenance_workbook(
    path: Path,
    all_rows: list[tuple[dict[str, str], str, str]],
    maintenance_rows: list[tuple[dict[str, str], str]],
    year: int,
    month: int,
    last_day: int,
) -> None:
    wb = Workbook()
    title = f"日立市役所様 本庁及び出先機関 各種賃貸借契約に対する保守対応＜{fullwidth(year)}年{fullwidth(month)}月度＞月次作業報告書"
    period = f"{fullwidth(year)}年{fullwidth(month)}月1日～{fullwidth(year)}年{fullwidth(month)}月{fullwidth(last_day)}日"
    setup_summary_sheet(wb.active, title, period, len(maintenance_rows), maintenance_remarks(maintenance_rows), True)
    data = wb.create_sheet("データ")
    data.append(MAINTENANCE_HEADERS)
    for index, (row, _flag) in enumerate(maintenance_rows, start=1):
        data.append(maintenance_row(row, index))
    style_report_sheet(
        data,
        [6, 12, 13, 13, 12, 22, 18, 16, 18, 18, 46, 46, 56],
        left_aligned_columns={11, 12, 13},
        border_style="medium",
    )
    for row_number, (_row, flag) in enumerate(maintenance_rows, start=2):
        mark_output_row(data, row_number, 1, len(MAINTENANCE_HEADERS), flag)
    raw = wb.create_sheet("元データ２")
    for _ in range(1):
        raw.append([])
    raw.append(["", *MAINTENANCE_HEADERS])
    helpdesk_fill = PatternFill("solid", fgColor="DDEBF7")
    for index, (row, classification, flag) in enumerate(all_rows, start=1):
        raw.append(["", *maintenance_row(row, index)])
        if classification == "helpdesk":
            for cell in raw[raw.max_row][1:]:
                cell.fill = copy(helpdesk_fill)
    style_report_sheet(raw, [6, 12, 13, 13, 12, 22, 18, 16, 18, 18, 46, 46, 56], header_row=2, start_col=2, data_start_row=3)
    for row_number, (_row, _classification, flag) in enumerate(all_rows, start=3):
        mark_output_row(raw, row_number, 2, len(MAINTENANCE_HEADERS), flag)
    wb.save(path)


def generate_reports(issues_csv: Path, output_dir: Path | None = None) -> dict[str, object]:
    if not issues_csv.exists():
        raise FileNotFoundError(f"CSV file not found: {issues_csv}")
    source_header, source_rows, encoding = read_csv_rows(issues_csv)
    rows = sorted(rows_as_dicts(source_header, source_rows), key=sort_key)
    classified = [item for row in rows for item in split_work_items(row)]
    helpdesk_rows = [(row, flag) for row, classification, flag in classified if classification == "helpdesk"]
    maintenance_rows = [(row, flag) for row, classification, flag in classified if classification == "maintenance"]
    year, month, last_day = previous_month()
    helpdesk_name, maintenance_name = report_filenames(year, month)
    output_dir = output_dir or issues_csv.parent
    output_dir.mkdir(parents=True, exist_ok=True)
    helpdesk_path = output_dir / helpdesk_name
    maintenance_path = output_dir / maintenance_name
    write_helpdesk_workbook(helpdesk_path, helpdesk_rows, year, month, last_day)
    write_maintenance_workbook(maintenance_path, classified, maintenance_rows, year, month, last_day)
    return {
        "issues_csv": issues_csv,
        "encoding": encoding,
        "total": len(rows),
        "helpdesk_count": len(helpdesk_rows),
        "maintenance_count": len(maintenance_rows),
        "helpdesk_output": helpdesk_path,
        "maintenance_output": maintenance_path,
    }


def run_gui() -> int:
    try:
        import tkinter as tk
        from tkinter import filedialog, messagebox
    except Exception as exc:
        print(f"GUI could not be started: {exc}", file=sys.stderr)
        return 1

    root = tk.Tk()
    root.withdraw()
    root.update()

    issues_file = filedialog.askopenfilename(
        title="issues.csv を選択してください",
        filetypes=(("CSV files", "*.csv"), ("All files", "*.*")),
    )
    if not issues_file:
        return 0

    desktop_dir = Path.home() / "Desktop"
    initial_dir = desktop_dir if desktop_dir.exists() else Path(issues_file).parent
    try:
        # The outputs are intentionally saved next to the selected issues.csv.
        result = generate_reports(Path(issues_file), output_dir=Path(issues_file).parent)
    except Exception as exc:
        messagebox.showerror("Report generation failed", str(exc))
        return 1

    messagebox.showinfo(
        "Report generation completed",
        "\n".join(
            [
                "Excel files were created successfully.",
                "",
                f"Total issues: {result['total']}",
                f"Help desk: {result['helpdesk_count']}",
                f"Maintenance: {result['maintenance_count']}",
                "",
                "Saved in the same folder as the selected CSV:",
                f"1) {result['helpdesk_output']}",
                f"2) {result['maintenance_output']}",
            ]
        ),
    )
    return 0


def main() -> int:
    args = parse_args()
    if args.gui or args.issues_csv is None:
        return run_gui()

    issues_csv = Path(args.issues_csv)

    try:
        result = generate_reports(issues_csv, output_dir=Path(args.output_dir) if args.output_dir else None)
    except Exception as exc:
        print(str(exc), file=sys.stderr)
        return 1

    print(f"Read CSV: {result['issues_csv']}")
    print(f"Detected encoding: {result['encoding']}")
    print(f"Total issues: {result['total']}")
    print(f"Help desk: {result['helpdesk_count']} -> {result['helpdesk_output']}")
    print(f"Maintenance: {result['maintenance_count']} -> {result['maintenance_output']}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
