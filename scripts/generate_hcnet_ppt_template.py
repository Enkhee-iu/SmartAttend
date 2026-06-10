#!/usr/bin/env python3
"""Generate HCNET-inspired PowerPoint template from hcnet.co.jp brand colors."""

from __future__ import annotations

import os
from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import PP_ALIGN
from pptx.util import Inches, Pt

COLORS = {
    "primary": RGBColor(0x00, 0x8C, 0x41),
    "accent": RGBColor(0x00, 0xC0, 0x5B),
    "dark": RGBColor(0x24, 0x59, 0x3D),
    "light_bg": RGBColor(0xE5, 0xF9, 0xEE),
    "text": RGBColor(0x33, 0x33, 0x33),
    "text_light": RGBColor(0x66, 0x66, 0x66),
    "bg": RGBColor(0xF8, 0xF8, 0xF8),
    "white": RGBColor(0xFF, 0xFF, 0xFF),
}

FONT_JP = "Yu Gothic"
SLIDE_W = Inches(13.333)
SLIDE_H = Inches(7.5)


def set_font(run, size: int, bold: bool = False, color=None, name: str = FONT_JP) -> None:
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.name = name
    if color:
        run.font.color.rgb = color


def add_rect(slide, left, top, width, height, fill_color, line_color=None):
    shape = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, left, top, width, height)
    shape.fill.solid()
    shape.fill.fore_color.rgb = fill_color
    if line_color:
        shape.line.color.rgb = line_color
    else:
        shape.line.fill.background()
    return shape


def add_textbox(slide, left, top, width, height, text, size=18, bold=False, color=None, align=PP_ALIGN.LEFT):
    if color is None:
        color = COLORS["text"]
    tb = slide.shapes.add_textbox(left, top, width, height)
    tf = tb.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.alignment = align
    run = p.add_run()
    run.text = text
    set_font(run, size, bold, color)
    return tb


def add_header_bar(slide, title_text: str, subtitle: str | None = None) -> None:
    add_rect(slide, Inches(0), Inches(0), SLIDE_W, Inches(0.12), COLORS["accent"])
    add_rect(slide, Inches(0), Inches(0.12), SLIDE_W, Inches(0.95), COLORS["primary"])
    add_textbox(slide, Inches(0.6), Inches(0.25), Inches(10), Inches(0.5), title_text, 24, True, COLORS["white"])
    if subtitle:
        add_textbox(slide, Inches(0.6), Inches(0.65), Inches(10), Inches(0.35), subtitle, 12, False, COLORS["light_bg"])
    add_rect(slide, Inches(0), Inches(7.2), SLIDE_W, Inches(0.3), COLORS["dark"])


def add_footer(slide, page_num: int | None = None, company: str = "エイチ・シー・ネットワークス株式会社") -> None:
    add_textbox(slide, Inches(0.6), Inches(7.05), Inches(8), Inches(0.25), company, 9, False, COLORS["text_light"])
    if page_num:
        add_textbox(slide, Inches(12), Inches(7.05), Inches(0.8), Inches(0.25), str(page_num), 9, False, COLORS["text_light"], PP_ALIGN.RIGHT)


def build_template() -> Presentation:
    prs = Presentation()
    prs.slide_width = SLIDE_W
    prs.slide_height = SLIDE_H
    blank = prs.slide_layouts[6]

    slide = prs.slides.add_slide(blank)
    add_rect(slide, Inches(0), Inches(0), SLIDE_W, SLIDE_H, COLORS["white"])
    add_rect(slide, Inches(0), Inches(0), Inches(0.35), SLIDE_H, COLORS["primary"])
    add_rect(slide, Inches(0), Inches(2.8), SLIDE_W, Inches(2.2), COLORS["primary"])
    add_rect(slide, Inches(0), Inches(2.75), SLIDE_W, Inches(0.08), COLORS["accent"])
    add_textbox(slide, Inches(1), Inches(3.0), Inches(11), Inches(0.8), "プレゼンテーションタイトル", 36, True, COLORS["white"])
    add_textbox(slide, Inches(1), Inches(3.85), Inches(11), Inches(0.5), "サブタイトル / 部署名 / 日付", 18, False, COLORS["light_bg"])
    add_textbox(slide, Inches(1), Inches(6.2), Inches(8), Inches(0.4), "エイチ・シー・ネットワークス株式会社", 14, False, COLORS["primary"])
    add_textbox(slide, Inches(1), Inches(6.55), Inches(8), Inches(0.35), "ネットワークのトータルソリューション", 11, False, COLORS["text_light"])

    slide = prs.slides.add_slide(blank)
    add_rect(slide, Inches(0), Inches(0), SLIDE_W, SLIDE_H, COLORS["bg"])
    add_rect(slide, Inches(0), Inches(0), SLIDE_W, Inches(0.12), COLORS["accent"])
    add_rect(slide, Inches(0), Inches(2.5), Inches(8), Inches(2.5), COLORS["primary"])
    add_textbox(slide, Inches(0.8), Inches(3.0), Inches(7), Inches(0.9), "01", 48, True, COLORS["accent"])
    add_textbox(slide, Inches(0.8), Inches(3.7), Inches(7), Inches(0.8), "セクションタイトル", 32, True, COLORS["white"])
    add_textbox(slide, Inches(8.5), Inches(3.2), Inches(4), Inches(1.5), "Section\nSubtitle", 14, False, COLORS["text_light"])
    add_footer(slide, 2)

    slide = prs.slides.add_slide(blank)
    add_rect(slide, Inches(0), Inches(0), SLIDE_W, SLIDE_H, COLORS["white"])
    add_header_bar(slide, "コンテンツスライド", "Content Slide")
    items = [
        "統合ITインフラソリューションの提案",
        "ネットワーク・セキュリティ・無線LAN",
        "構築から保守サービスまでワンストップ対応",
        "高い品質（High quality）と信頼（Confidence）",
    ]
    tb = slide.shapes.add_textbox(Inches(0.9), Inches(1.5), Inches(11), Inches(5))
    tf = tb.text_frame
    tf.word_wrap = True
    for i, item in enumerate(items):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.space_after = Pt(16)
        run = p.add_run()
        run.text = f"●  {item}"
        set_font(run, 20, False, COLORS["text"])
    add_rect(slide, Inches(0.6), Inches(1.35), Inches(0.08), Inches(4.5), COLORS["accent"])
    add_footer(slide, 3)

    slide = prs.slides.add_slide(blank)
    add_rect(slide, Inches(0), Inches(0), SLIDE_W, SLIDE_H, COLORS["white"])
    add_header_bar(slide, "2カラムレイアウト", "Two Column Layout")
    add_rect(slide, Inches(0.6), Inches(1.4), Inches(5.8), Inches(5.3), COLORS["light_bg"])
    add_rect(slide, Inches(0.6), Inches(1.4), Inches(5.8), Inches(0.06), COLORS["primary"])
    add_textbox(slide, Inches(0.9), Inches(1.6), Inches(5.2), Inches(0.5), "左カラム", 20, True, COLORS["primary"])
    add_textbox(slide, Inches(0.9), Inches(2.2), Inches(5.2), Inches(4), "製品・サービス\n\n● Adapterシリーズ\n● ネットワーク\n● 無線LAN\n● セキュリティ", 16)
    add_rect(slide, Inches(6.9), Inches(1.4), Inches(5.8), Inches(5.3), COLORS["bg"])
    add_rect(slide, Inches(6.9), Inches(1.4), Inches(5.8), Inches(0.06), COLORS["accent"])
    add_textbox(slide, Inches(7.2), Inches(1.6), Inches(5.2), Inches(0.5), "右カラム", 20, True, COLORS["primary"])
    add_textbox(slide, Inches(7.2), Inches(2.2), Inches(5.2), Inches(4), "ソリューション\n\n● 企業向けDX\n● キャンパスネットワーク\n● 医療ネットワーク\n● 社会インフラ", 16)
    add_footer(slide, 4)

    slide = prs.slides.add_slide(blank)
    add_rect(slide, Inches(0), Inches(0), SLIDE_W, SLIDE_H, COLORS["white"])
    add_header_bar(slide, "画像＋テキスト", "Image + Text")
    add_rect(slide, Inches(0.6), Inches(1.4), Inches(6), Inches(5.3), COLORS["bg"])
    add_textbox(slide, Inches(1.5), Inches(3.5), Inches(4), Inches(0.5), "[ 画像をここに配置 ]", 16, False, COLORS["text_light"], PP_ALIGN.CENTER)
    shape = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0.6), Inches(1.4), Inches(6), Inches(5.3))
    shape.fill.background()
    shape.line.color.rgb = COLORS["primary"]
    shape.line.width = Pt(1)
    add_rect(slide, Inches(7), Inches(1.4), Inches(5.7), Inches(0.06), COLORS["primary"])
    add_textbox(slide, Inches(7), Inches(1.6), Inches(5.5), Inches(0.5), "ポイント", 22, True, COLORS["primary"])
    for i, pt in enumerate(["HiVAS運用サービス", "IT Asset管理", "SSEセキュリティ", "IEEE802.11ah対応"]):
        y = 2.3 + i * 0.9
        add_rect(slide, Inches(7), Inches(y), Inches(0.35), Inches(0.35), COLORS["accent"])
        add_textbox(slide, Inches(7.5), Inches(y - 0.05), Inches(5), Inches(0.4), pt, 16)
    add_footer(slide, 5)

    slide = prs.slides.add_slide(blank)
    add_rect(slide, Inches(0), Inches(0), SLIDE_W, SLIDE_H, COLORS["white"])
    add_header_bar(slide, "表・データ", "Table / Data")
    table_shape = slide.shapes.add_table(5, 4, Inches(0.6), Inches(1.5), Inches(12), Inches(4.5))
    table = table_shape.table
    headers = ["項目", "内容", "担当", "ステータス"]
    data = [
        ["製品A", "ネットワーク機器", "営業部", "対応中"],
        ["製品B", "セキュリティ", "技術部", "完了"],
        ["製品C", "無線LAN", "営業部", "計画中"],
        ["製品D", "サービス", "サポート", "対応中"],
    ]
    for j, header in enumerate(headers):
        cell = table.cell(0, j)
        cell.text = header
        cell.fill.solid()
        cell.fill.fore_color.rgb = COLORS["primary"]
        for paragraph in cell.text_frame.paragraphs:
            for run in paragraph.runs:
                set_font(run, 14, True, COLORS["white"])
    for i, row in enumerate(data):
        for j, value in enumerate(row):
            cell = table.cell(i + 1, j)
            cell.text = value
            if i % 2 == 0:
                cell.fill.solid()
                cell.fill.fore_color.rgb = COLORS["light_bg"]
            for paragraph in cell.text_frame.paragraphs:
                for run in paragraph.runs:
                    set_font(run, 13, False, COLORS["text"])
    add_footer(slide, 6)

    slide = prs.slides.add_slide(blank)
    add_rect(slide, Inches(0), Inches(0), SLIDE_W, SLIDE_H, COLORS["primary"])
    add_rect(slide, Inches(0), Inches(0), SLIDE_W, Inches(0.12), COLORS["accent"])
    add_textbox(slide, Inches(1.5), Inches(2.5), Inches(10), Inches(1.5), '"高い品質と信頼できる\nネットワークを基本に"', 32, True, COLORS["white"], PP_ALIGN.CENTER)
    add_textbox(slide, Inches(1.5), Inches(4.5), Inches(10), Inches(0.5), "— HCNET Vision —", 16, False, COLORS["light_bg"], PP_ALIGN.CENTER)
    add_footer(slide, 7)

    slide = prs.slides.add_slide(blank)
    add_rect(slide, Inches(0), Inches(0), SLIDE_W, SLIDE_H, COLORS["white"])
    add_rect(slide, Inches(0), Inches(3.0), SLIDE_W, Inches(1.5), COLORS["primary"])
    add_rect(slide, Inches(0), Inches(2.95), SLIDE_W, Inches(0.08), COLORS["accent"])
    add_textbox(slide, Inches(0), Inches(3.2), SLIDE_W, Inches(0.8), "Thank You", 40, True, COLORS["white"], PP_ALIGN.CENTER)
    add_textbox(slide, Inches(0), Inches(4.0), SLIDE_W, Inches(0.5), "ご清聴ありがとうございました", 18, False, COLORS["light_bg"], PP_ALIGN.CENTER)
    add_textbox(slide, Inches(0), Inches(5.5), SLIDE_W, Inches(0.4), "www.hcnet.co.jp  |  お問い合わせ", 12, False, COLORS["text_light"], PP_ALIGN.CENTER)

    slide = prs.slides.add_slide(blank)
    add_rect(slide, Inches(0), Inches(0), SLIDE_W, SLIDE_H, COLORS["bg"])
    add_header_bar(slide, "テンプレート使用ガイド", "Template Guide")
    guide = [
        "スライド1: タイトル — 表紙",
        "スライド2: セクション区切り",
        "スライド3: 箇条書きコンテンツ",
        "スライド4: 2カラムレイアウト",
        "スライド5: 画像＋テキスト",
        "スライド6: 表・データ",
        "スライド7: 引用・ハイライト",
        "スライド8: 終了（Thank You）",
    ]
    tb = slide.shapes.add_textbox(Inches(0.9), Inches(1.5), Inches(5.5), Inches(5))
    tf = tb.text_frame
    for i, line in enumerate(guide):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.space_after = Pt(10)
        run = p.add_run()
        run.text = line
        set_font(run, 15, False, COLORS["text"])
    colors_info = [
        ("Primary Green", "#008C41", COLORS["primary"]),
        ("Accent Green", "#00C05B", COLORS["accent"]),
        ("Dark Green", "#24593D", COLORS["dark"]),
        ("Light BG", "#E5F9EE", COLORS["light_bg"]),
        ("Text", "#333333", COLORS["text"]),
        ("Background", "#F8F8F8", COLORS["bg"]),
    ]
    add_textbox(slide, Inches(7), Inches(1.5), Inches(5), Inches(0.4), "カラーパレット（HCNET公式サイトより）", 16, True, COLORS["primary"])
    for i, (name, hex_code, rgb) in enumerate(colors_info):
        y = 2.0 + i * 0.75
        add_rect(slide, Inches(7), Inches(y), Inches(0.6), Inches(0.5), rgb)
        add_textbox(slide, Inches(7.8), Inches(y + 0.05), Inches(4.5), Inches(0.4), f"{name}  {hex_code}", 13)
    add_footer(slide, 9)

    return prs


def main() -> None:
    root = Path(__file__).resolve().parents[1]
    out_dir = root / "docs" / "templates"
    out_dir.mkdir(parents=True, exist_ok=True)
    out_path = out_dir / "HCNET_PowerPoint_Template.pptx"
    build_template().save(out_path)
    print(f"Saved: {out_path}")


if __name__ == "__main__":
    main()
