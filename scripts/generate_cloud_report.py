"""Generate HCNET private cloud monthly reports from graph data workbooks."""

from __future__ import annotations

import re
import shutil
import sys
import tempfile
import xml.etree.ElementTree as ET
from datetime import datetime, timezone
from pathlib import Path
from zipfile import ZIP_DEFLATED, ZipFile

from openpyxl import load_workbook

TEMPLATE_RELATIVE_PATH = Path("templates") / "hcnet_cloud_monthly_template.xlsx"


SHEET_MAP = {
    "cluster": {
        "target_prefix": "Cluster",
        "target_name": "Cluster_{yyyymm}",
        "columns": [
            ("Name", "Timestamp"),
            ("Memory Usage (%)", "Memory Usage (%)（表示）"),
            ("CPU Usage (%)", "CPU Usage (%)（表示）"),
        ],
    },
    "hostCPU": {
        "target_prefix": "host_CPU",
        "target_name": "host_CPU_{yyyymm}",
        "columns": [
            ("Name", "Timestamp"),
            ("ntnx06-ahv01", "ntnx06-ahv01（表示）"),
            ("ntnx06-ahv02", "ntnx06-ahv02（表示）"),
            ("ntnx06-ahv03", "ntnx06-ahv03（表示）"),
            ("ntnx06-ahv04", "ntnx06-ahv04（表示）"),
            ("ntnx06-ahv05", "ntnx06-ahv05（表示）"),
        ],
    },
    "hostMemory": {
        "target_prefix": "host_memory",
        "target_name": "host_memory_{yyyymm}",
        "columns": [
            ("Name", "Timestamp"),
            ("ntnx06-ahv01", "ntnx06-ahv01（表示）"),
            ("ntnx06-ahv02", "ntnx06-ahv02（表示）"),
            ("ntnx06-ahv03", "ntnx06-ahv03（表示）"),
            ("ntnx06-ahv04", "ntnx06-ahv04（表示）"),
            ("ntnx06-ahv05", "ntnx06-ahv05（表示）"),
        ],
    },
    "storage": {
        "target_prefix": "storage",
        "target_name": "storage_{yyyymm}",
        "columns": [
            ("Name", "Timestamp"),
            ("データストア利用容量", "TB表示(/1024)"),
            ("データストア利用率", "利用率(全体を181.4TBで算出)"),
        ],
    },
}


def infer_report_yyyymm(*paths: Path) -> str:
    for path in paths:
        match = re.search(r"(20\d{4})", path.name)
        if match:
            return match.group(1)
    raise ValueError("ファイル名からレポート年月(YYYYMM)を取得できません。")


def resource_path(relative_path: Path) -> Path:
    base_path = Path(getattr(sys, "_MEIPASS", Path(__file__).resolve().parents[1]))
    return base_path / relative_path


def infer_report_yyyymm_from_workbook(path: Path) -> str:
    wb = load_workbook(path, read_only=True, data_only=True)
    for sheet_name in wb.sheetnames:
        match = re.fullmatch(r"(20\d{2})-(\d{2})(?:\(1\))?", sheet_name)
        if not match:
            continue
        year = int(match.group(1))
        month = int(match.group(2)) + 1
        if month == 13:
            month = 1
            year += 1
        return f"{year}{month:02d}"
    raise ValueError("ファイル名またはシート名からレポート年月(YYYYMM)を取得できません。")


def previous_month_yyyymm(report_yyyymm: str) -> str:
    year = int(report_yyyymm[:4])
    month = int(report_yyyymm[4:])
    month -= 1
    if month == 0:
        month = 12
        year -= 1
    return f"{year}{month:02d}"


def yyyymm_to_sheet_month(yyyymm: str) -> str:
    return f"{yyyymm[:4]}-{yyyymm[4:]}"


def last_day(year: int, month: int) -> int:
    import calendar

    return calendar.monthrange(year, month)[1]


def timestamp_to_datetime(value) -> datetime:
    if isinstance(value, datetime):
        return value.replace(tzinfo=None)
    if value is None:
        raise ValueError("Timestamp is empty")
    number = int(value)
    # Graph exports use Unix timestamps in microseconds.
    if number > 10_000_000_000:
        number = number / 1_000_000
    return datetime.fromtimestamp(number, timezone.utc).replace(tzinfo=None)


def headers(ws) -> dict[str, int]:
    return {str(ws.cell(1, col).value): col for col in range(1, ws.max_column + 1)}


def find_sheet_by_prefix(wb, prefix: str):
    for ws in wb.worksheets:
        if ws.title.startswith(prefix + "_"):
            return ws
    raise KeyError(f"Sheet starting with {prefix}_ was not found")


def set_cell_values(ws, values: list[list[object]]) -> None:
    for row_idx, row in enumerate(values, start=1):
        for col_idx, value in enumerate(row, start=1):
            ws.cell(row_idx, col_idx).value = value


def clear_sheet_values(ws, max_columns: int) -> None:
    for row in ws.iter_rows(min_row=1, max_row=ws.max_row, min_col=1, max_col=max(ws.max_column, max_columns)):
        for cell in row:
            cell.value = None


def source_rows(ws, column_pairs: list[tuple[str, str]], data_yyyymm: str) -> list[list[object]]:
    src_headers = headers(ws)
    missing = [source for _target, source in column_pairs if source not in src_headers]
    if missing:
        raise KeyError(f"{ws.title} に必要な列がありません: {', '.join(missing)}")

    year = int(data_yyyymm[:4])
    month = int(data_yyyymm[4:])
    start = datetime(year, month, 1)
    end_month = month + 1
    end_year = year
    if end_month == 13:
        end_month = 1
        end_year += 1
    end = datetime(end_year, end_month, 1)

    result: list[list[object]] = []
    timestamp_col = src_headers["Timestamp"]
    for row_idx in range(2, ws.max_row + 1):
        raw_timestamp = ws.cell(row_idx, timestamp_col).value
        if raw_timestamp in (None, ""):
            continue
        converted = timestamp_to_datetime(raw_timestamp)
        if not (start <= converted < end):
            continue
        row_values = []
        for _target, source in column_pairs:
            if source == "Timestamp":
                row_values.append(converted)
            else:
                row_values.append(ws.cell(row_idx, src_headers[source]).value)
        result.append(row_values)

    result.sort(key=lambda row: row[0])
    return result


def update_summary_sheet(ws, report_yyyymm: str, data_yyyymm: str) -> None:
    report_year = int(report_yyyymm[:4])
    report_month = int(report_yyyymm[4:])
    data_year = int(data_yyyymm[:4])
    data_month = int(data_yyyymm[4:])
    data_last_day = last_day(data_year, data_month)

    ws["J2"] = f"{report_year}年 {report_month}月 20日"

    title_text = f"HCNETプライベートクラウド仮想環境＜{data_year}年 {data_month}月度＞月次報告書"
    period_text = f"{data_year}年 {data_month}月 1日～{data_year}年 {data_month}月 {data_last_day}日"

    for row in ws.iter_rows():
        for cell in row:
            if not isinstance(cell.value, str):
                continue
            if "HCNETプライベートクラウド仮想環境" in cell.value and "月次報告書" in cell.value:
                cell.value = title_text
            elif cell.value.startswith("20") and "月 1日～" in cell.value:
                cell.value = period_text


def rename_overview_sheets(wb, data_sheet_month: str) -> None:
    month_sheets = [ws for ws in wb.worksheets if re.fullmatch(r"20\d{2}-\d{2}(?:\\(1\\))?", ws.title)]
    base_sheet = next((ws for ws in month_sheets if not ws.title.endswith("(1)")), None)
    resource_sheet = next((ws for ws in month_sheets if ws.title.endswith("(1)")), None)
    if base_sheet is not None:
        base_sheet.title = data_sheet_month
    if resource_sheet is not None:
        resource_sheet.title = f"{data_sheet_month}(1)"


def populate_target_sheet(target_ws, source_ws, config: dict, data_yyyymm: str) -> None:
    column_pairs = config["columns"]
    rows = source_rows(source_ws, column_pairs, data_yyyymm)
    values = [[target for target, _source in column_pairs], *rows]

    clear_sheet_values(target_ws, len(column_pairs))
    set_cell_values(target_ws, values)
    for row in target_ws.iter_rows(min_row=2, max_row=target_ws.max_row, min_col=1, max_col=1):
        row[0].number_format = "yyyy-mm-dd hh:mm:ss"


def add_content_type_override(root, part_name: str, content_type: str) -> None:
    namespace = "http://schemas.openxmlformats.org/package/2006/content-types"
    existing = {
        element.attrib.get("PartName")
        for element in root.findall(f"{{{namespace}}}Override")
    }
    if part_name in existing:
        return
    ET.SubElement(root, f"{{{namespace}}}Override", PartName=part_name, ContentType=content_type)


def add_workbook_connection_relationship(root) -> None:
    namespace = "http://schemas.openxmlformats.org/package/2006/relationships"
    connection_type = "http://schemas.openxmlformats.org/officeDocument/2006/relationships/connections"
    for element in root.findall(f"{{{namespace}}}Relationship"):
        if element.attrib.get("Type") == connection_type:
            return
    used_ids = {
        element.attrib.get("Id", "")
        for element in root.findall(f"{{{namespace}}}Relationship")
    }
    next_index = 1
    while f"rId{next_index}" in used_ids:
        next_index += 1
    ET.SubElement(
        root,
        f"{{{namespace}}}Relationship",
        Id=f"rId{next_index}",
        Type=connection_type,
        Target="connections.xml",
    )


def restore_query_table_metadata(template_path: Path, workbook_path: Path) -> None:
    """Restore Excel query-table metadata removed by openpyxl.

    The HCNET workbook contains Excel tables backed by queryTable/external data
    range metadata. openpyxl preserves the visible cells and charts, but it drops
    these unsupported package parts when saving. Excel then repairs the workbook
    by deleting the external data ranges. Restoring these parts keeps the workbook
    opening cleanly and preserves table behavior.
    """
    with ZipFile(template_path) as template_zip, ZipFile(workbook_path) as generated_zip:
        template_names = set(template_zip.namelist())
        generated_names = set(generated_zip.namelist())

        restore_names = [
            name
            for name in template_names
            if name.startswith("xl/queryTables/")
            or name.startswith("xl/tables/_rels/")
            or name == "xl/connections.xml"
        ]

        content_root = ET.fromstring(generated_zip.read("[Content_Types].xml"))
        if "xl/connections.xml" in restore_names:
            add_content_type_override(
                content_root,
                "/xl/connections.xml",
                "application/vnd.openxmlformats-officedocument.spreadsheetml.connections+xml",
            )
        for name in restore_names:
            if name.startswith("xl/queryTables/") and name.endswith(".xml"):
                add_content_type_override(
                    content_root,
                    "/" + name,
                    "application/vnd.openxmlformats-officedocument.spreadsheetml.queryTable+xml",
                )

        workbook_rels_root = ET.fromstring(generated_zip.read("xl/_rels/workbook.xml.rels"))
        if "xl/connections.xml" in restore_names:
            add_workbook_connection_relationship(workbook_rels_root)

        with tempfile.NamedTemporaryFile(delete=False, suffix=".xlsx") as tmp_file:
            tmp_path = Path(tmp_file.name)

        try:
            with ZipFile(tmp_path, "w", ZIP_DEFLATED) as output_zip:
                for name in generated_zip.namelist():
                    if name in {"[Content_Types].xml", "xl/_rels/workbook.xml.rels"}:
                        continue
                    if name in restore_names:
                        continue
                    output_zip.writestr(name, generated_zip.read(name))

                output_zip.writestr(
                    "[Content_Types].xml",
                    ET.tostring(content_root, encoding="utf-8", xml_declaration=True),
                )
                output_zip.writestr(
                    "xl/_rels/workbook.xml.rels",
                    ET.tostring(workbook_rels_root, encoding="utf-8", xml_declaration=True),
                )
                for name in restore_names:
                    output_zip.writestr(name, template_zip.read(name))

            shutil.move(tmp_path, workbook_path)
        finally:
            if tmp_path.exists():
                tmp_path.unlink()


def output_filename(report_yyyymm: str) -> str:
    return f"HCNETプライベートクラウド月次報告_グラフ込み({report_yyyymm})原本_20260424.xlsx"


def generate_cloud_report(
    graph_template_path: Path,
    output_path: Path | None = None,
    template_path: Path | None = None,
) -> Path:
    try:
        report_yyyymm = infer_report_yyyymm(graph_template_path)
    except ValueError:
        report_yyyymm = infer_report_yyyymm_from_workbook(graph_template_path)
    data_yyyymm = previous_month_yyyymm(report_yyyymm)
    data_sheet_month = yyyymm_to_sheet_month(data_yyyymm)

    template_path = template_path or resource_path(TEMPLATE_RELATIVE_PATH)
    if not template_path.exists():
        raise FileNotFoundError(f"月次報告書テンプレートが見つかりません: {template_path}")

    output_path = output_path or graph_template_path.with_name(output_filename(report_yyyymm))
    shutil.copy2(template_path, output_path)

    target_wb = load_workbook(output_path)
    source_wb = load_workbook(graph_template_path, data_only=True)

    rename_overview_sheets(target_wb, data_sheet_month)
    if data_sheet_month in target_wb.sheetnames:
        update_summary_sheet(target_wb[data_sheet_month], report_yyyymm, data_yyyymm)

    for source_sheet_name, config in SHEET_MAP.items():
        if source_sheet_name not in source_wb.sheetnames:
            raise KeyError(f"グラフひな形に {source_sheet_name} シートがありません。")
        target_ws = find_sheet_by_prefix(target_wb, config["target_prefix"])
        target_ws.title = config["target_name"].format(yyyymm=data_yyyymm)
        populate_target_sheet(target_ws, source_wb[source_sheet_name], config, data_yyyymm)

    target_wb.save(output_path)
    restore_query_table_metadata(template_path, output_path)
    return output_path


def run_gui(parent=None) -> None:
    try:
        from tkinter import filedialog, messagebox
    except Exception:
        return

    graph_template = filedialog.askopenfilename(
        parent=parent,
        title="グラフひな形 Excel を選択してください",
        filetypes=(("Excel files", "*.xlsx"), ("All files", "*.*")),
    )
    if not graph_template:
        return

    try:
        output_path = generate_cloud_report(Path(graph_template))
    except Exception as exc:
        messagebox.showerror("HCNETクラウド月次報告書", f"作成に失敗しました。\n\n{exc}", parent=parent)
        return

    messagebox.showinfo("HCNETクラウド月次報告書", f"作成が完了しました。\n\n{output_path}", parent=parent)
